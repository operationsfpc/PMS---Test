import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_date_wise_user_manual_changes_docx(output_path: Path):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("FPC Placement Management System — September 2026 Release Changelog")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(150, 150, 150)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential — Placement Operations Documentation — September 04 to September 11, 2026")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(150, 150, 150)

    # Palette
    PRIMARY = RGBColor(61, 55, 119)     # Deep Indigo #3D3777
    SECONDARY = RGBColor(217, 119, 6)   # Amber #D97706
    TEXT = RGBColor(40, 40, 40)         # Charcoal
    MUTED = RGBColor(100, 116, 139)     # Slate

    # Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_brand = p_title.add_run("FACE Prep Campus — Placement Management System\n")
    r_brand.font.name = "Segoe UI"
    r_brand.font.size = Pt(12)
    r_brand.font.bold = True
    r_brand.font.color.rgb = SECONDARY

    r_main = p_title.add_run("User Manual Updates & Release Changelog")
    r_main.font.name = "Segoe UI"
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Date-wise operational record of system enhancements, user manual adjustments, and bug fixes from September 04 to September 11, 2026.")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = MUTED

    def add_release_heading(date_str, title_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        
        r_badge = p.add_run(f"[{date_str}] ")
        r_badge.font.name = "Segoe UI"
        r_badge.font.size = Pt(12)
        r_badge.font.bold = True
        r_badge.font.color.rgb = SECONDARY
        
        run = p.add_run(title_str)
        run.font.name = "Segoe UI"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = PRIMARY
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(bold_prefix)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = PRIMARY
        r2 = p.add_run(text)
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT
        return p

    # -------------------------------------------------------------
    # 1. September 11, 2026
    # -------------------------------------------------------------
    add_release_heading("September 11, 2026", "SRF Dropdown Synchronization, Conditional Field Sanitization & Label Validation")
    add_bullet("Dropdown Value Replacement & State Sync: ",
               "When modifying any dropdown selection in the Student Registration Form (such as school board, marks scale, or degree fork), previous values are cleanly replaced. Form state stays completely synchronized with what the student sees on screen.")
    add_bullet("Automatic Obsolete Field Sanitization: ",
               "Switching away from conditional options (e.g. from State Board to CBSE, or from PG to UG) immediately wipes obsolete child fields (State Board state, Other board name, Cambridge/Other grades, completed UG degree details, and consolidated marksheets) from both the active form state and saved drafts. This eliminates hidden validation blocks caused by lingering orphan data.")
    add_bullet("Immediate Scale Switch Revalidation: ",
               "Changing the College Marks Scale between CGPA (10-point scale) and Percentage immediately re-evaluates all semester figures in real time without requiring the student to click submit.")
    add_bullet("Field Label Validation Highlighting (Red Labels & 'Invalid' Badge): ",
               "Every required or invalid field now directly highlights its label in bold red (text-danger-700 font-semibold) with a clear, accessible 'Invalid' badge. Students can immediately see which exact field is missing or contains formatting errors.")
    add_bullet("Real-Time Error Clearing: ",
               "As soon as the user corrects an invalid input or selects a valid option, the red label styling, error badge, and red border clear instantly in real time.")
    add_bullet("Consolidated Error Messaging: ",
               "Removed duplicate alert nodes and standardized error rendering across all inputs, ensuring clean DOM structure and eliminating confusing repeated messages.")

    # -------------------------------------------------------------
    # 2. September 10, 2026
    # -------------------------------------------------------------
    add_release_heading("September 10, 2026", "Coordinator Verification Queue Enhancements & Database Arrears Sync (Migrations 0071–0074)")
    add_bullet("International Board Letter Grades Display: ",
               "In the CPC verification queue (/cpc/verification), the 10th and 12th marks columns now display letter grades (e.g., 'Grade: A*' or '91.4% (Grade: A*)') so coordinators can easily verify Cambridge (IGCSE/A-Levels) and Other board results without confusion.")
    add_bullet("Arrear History Discrepancy Resolution (2 vs 1): ",
               "Fixed a bug where a student who declared 2 arrears across degree semesters showed only 1 in the coordinator queue. The queue now accurately evaluates the true cumulative maximum across all declared semesters.")
    add_bullet("Automated Standing Arrears Synchronization: ",
               "The student's declared standing arrears from their latest semester are automatically synchronized into the core students table on form submission (via Migration 0073), ensuring recruiter drive eligibility filters always read current academic data.")
    add_bullet("School Marksheets Column Isolation: ",
               "The verification queue 'School Marksheets' column is strictly isolated to educational qualification documents (10th, 12th, Diploma, Consolidated UG). Resumes and certificates are separated into their own designated columns.")
    add_bullet("Active Marksheet Scan Deduplication: ",
               "When a student re-uploads a corrected document, only the single latest active scan is displayed, preventing broken links and duplicate clutter.")
    add_bullet("Staff & Mentor Password Management: ",
               "Added secure database password reset capabilities for campus coordinators and placement mentors (Migration 0071).")

    # -------------------------------------------------------------
    # 3. September 04, 2026
    # -------------------------------------------------------------
    add_release_heading("September 04, 2026", "Automated Email Notification System & Catalog Release")
    add_bullet("Automated Transactional Email Triggers: ",
               "Activated automated email dispatches for key placement milestones: SRF Approval Confirmation, SRF Rejection (with specific coordinator notes), New Placement Drive Announcements, Student Drive Registration Confirmation, and Shortlist Advancement Updates.")
    add_bullet("Official Email Notification Catalog: ",
               "Published comprehensive email catalog (docs/FACE_Prep_PMS_Email_Notification_Catalog.docx) detailing recipient rules, delivery triggers, dynamic template variables, and layout previews.")
    add_bullet("Resilient Dispatch Architecture: ",
               "Implemented decoupled email queue handling and background execution fallbacks so batch email operations never impede web interface responsiveness or core database transactions.")

    # -------------------------------------------------------------
    # 4. Chronological Summary Table (Sept 04 – Sept 11)
    # -------------------------------------------------------------
    p_tbl_heading = doc.add_paragraph()
    p_tbl_heading.paragraph_format.space_before = Pt(16)
    p_tbl_heading.paragraph_format.space_after = Pt(4)
    r_th = p_tbl_heading.add_run("Date-Wise Summary: September 04 to September 11, 2026")
    r_th.font.name = "Segoe UI"
    r_th.font.size = Pt(12)
    r_th.font.bold = True
    r_th.font.color.rgb = PRIMARY

    table_data = [
        ("Date", "Feature / Area", "Previous State (Problem)", "Updated Behavior (Proper Fix)"),
        ("2026-09-11", "Dropdown Replacement", "Changing dropdowns left old data or orphan fields behind.", "Dropdown choices cleanly replace previous values and reset child fields."),
        ("2026-09-11", "Conditional Sanitization", "Old state board states / PG entries persisted in drafts.", "Obsolete conditional values are automatically sanitized on live change and draft merge."),
        ("2026-09-11", "Field Label Highlighting", "Labels remained plain grey when inputs were empty or invalid.", "Field labels highlight in bold red (text-danger-700) with an 'Invalid' badge."),
        ("2026-09-11", "Real-Time Error Clearing", "Error indicators remained until full form submit.", "Label highlighting and invalid badge clear instantly as user types or picks value."),
        ("2026-09-11", "Scale Switch Validation", "Toggling CGPA/Percentage scale did not revalidate marks.", "Immediately re-evaluates all semester marks upon scale change in real time."),
        ("2026-09-10", "Verification Board Grades", "Cambridge/Other letter grades were invisible to CPCs.", "Displays letter grades (Grade: A*) alongside percentage in verification queue."),
        ("2026-09-10", "Arrear History Accuracy", "Queue showed 1 arrear history when student entered 2.", "Accurately calculates cumulative maximum arrear history across all semesters."),
        ("2026-09-10", "Database Arrears Sync", "Latest arrears were not synced to students table.", "Database trigger automatically synchronizes standing arrears on submission."),
        ("2026-09-10", "Marksheet Isolation", "Queue marksheet column was cluttered with resumes and certs.", "Shows strictly educational marksheets; resumes and certs in dedicated columns."),
        ("2026-09-04", "Email Notification Engine", "No automated email alerts sent to students or staff.", "Automated transactional emails for SRF approval, rejection, and drives."),
        ("2026-09-04", "Notification Catalog", "No standardized documentation for email templates.", "Published official email notification catalog with triggers and templates.")
    ]

    t = doc.add_table(rows=len(table_data), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)

    col_widths = [Inches(1.0), Inches(1.8), Inches(2.1), Inches(2.3)]

    for r_idx, row in enumerate(table_data):
        is_hdr = (r_idx == 0)
        bg = "3D3777" if is_hdr else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")

        for c_idx, text in enumerate(row):
            cell = t.rows[r_idx].cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1

            run = p.add_run(text)
            run.font.name = "Segoe UI"
            if is_hdr:
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.size = Pt(8.5)
                if c_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = SECONDARY
                elif c_idx == 1:
                    run.font.bold = True
                    run.font.color.rgb = PRIMARY
                else:
                    run.font.color.rgb = TEXT

    doc.save(str(output_path))
    print(f"Successfully generated clean date-wise docx (Sept 04-11) at {output_path}")

if __name__ == "__main__":
    out_file = Path("docs/USER_MANUAL_CHANGES.docx")
    build_date_wise_user_manual_changes_docx(out_file)
