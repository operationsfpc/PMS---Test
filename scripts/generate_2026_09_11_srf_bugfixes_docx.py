import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_docx(output_path: Path):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("FPC Placement Management System — Engineering Spec (2026-09-11)")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(150, 150, 150)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential — Placement Operations Documentation — September 11, 2026")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(150, 150, 150)

    PRIMARY = RGBColor(61, 55, 119)     # Deep Indigo #3D3777
    SECONDARY = RGBColor(217, 119, 6)   # Amber #D97706
    TEXT = RGBColor(40, 40, 40)         # Charcoal
    MUTED = RGBColor(100, 116, 139)     # Slate
    DANGER = RGBColor(185, 28, 28)      # Danger Red

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_brand = p_title.add_run("FACE Prep Campus — PMS\n")
    r_brand.font.name = "Segoe UI"
    r_brand.font.size = Pt(12)
    r_brand.font.bold = True
    r_brand.font.color.rgb = SECONDARY

    r_main = p_title.add_run("SRF Dropdown Synchronization & Label Validation Bug Fixes")
    r_main.font.name = "Segoe UI"
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Release Date: September 11, 2026 | Status: Verified & Tested (44/44 Tests Passing)")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = MUTED

    def add_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.font.name = "Segoe UI"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = PRIMARY
        return p

    def add_point(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(bold_prefix)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = PRIMARY
        r2 = p.add_run(text)
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT
        return p

    # Section 1: Overview
    add_heading("1. Problem Statement & Background")
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.15
    r_intro = p_intro.add_run(
        "Users reported that when filling fields in the Student Registration Form (SRF) using dropdown options "
        "and later modifying previously selected values, the old data was not being updated correctly. "
        "Obsolete conditional fields (such as state selections or other board names) remained persisted in the "
        "form state and saved drafts. Furthermore, field validation errors were only attached to input borders, "
        "leaving labels un-highlighted and making it difficult for students to spot exactly which fields were invalid."
    )
    r_intro.font.name = "Segoe UI"
    r_intro.font.size = Pt(9.5)
    r_intro.font.color.rgb = TEXT

    # Section 2: Key Fixes
    add_heading("2. Core Bug Fixes & Improvements Implemented")
    add_point("Dropdown Values Correctly Replace Previous Data: ",
              "Dedicated change handlers reset and replace all obsolete conditional values immediately when a new dropdown option is chosen.")
    add_point("Obsolete Conditional State Sanitization: ",
              "When switching boards (e.g. from State Board to CBSE), State Board State, Other Board Name, and Grade fields are automatically wiped from form state so stale data cannot fail validation or persist to database.")
    add_point("Programme Level Fork Cleanup: ",
              "Switching from Postgraduate (PG) back to Undergraduate (UG) immediately cleanses completed UG degree fields (ugDegree, ugCollege, ugBranch, ugAggregate, and ug_consolidated marksheet).")
    add_point("College Marks Scale Immediate Revalidation: ",
              "Changing the college marks scale between CGPA and Percentage immediately triggers validation on all entered semester marks, giving instant feedback without waiting for form submission.")
    add_point("Draft Sanitization on Resume/Merge: ",
              "Added sanitizeConditionalSrfValues() to ensure drafts saved under previous selections are stripped of stale conditional data upon being loaded.")
    add_point("Field Label Validation Highlighting: ",
              "Every Field, BoardSelect, and ScaleSelect label now dynamically highlights in bold danger red (text-danger-700 font-semibold) with an accessible 'Invalid' badge whenever invalid or empty on submit.")
    add_point("Real-Time Highlight Clearing: ",
              "As soon as the user enters valid data or selects an option, the label highlighting and 'Invalid' badge clear instantly without needing page reload.")
    add_point("DOM Error Consolidation: ",
              "Eliminated duplicate <ErrorText> tags across all fields, ensuring clean DOM structure and 100% test reliability in Vitest/Testing Library.")

    # Section 3: Summary Table
    add_heading("3. Verification Matrix & Before vs. After")

    table_data = [
        ("Component / Scenario", "Before (Buggy Behavior)", "After (Fixed & Synchronized)"),
        ("10th/12th Board Change", "Switching from State Board to CBSE kept state in form state, causing validation failure.", "Obsolete state and other board fields are cleared and revalidated immediately."),
        ("UG / PG Switch", "Switching to UG kept completed UG degree and marksheets from previous PG state.", "All completed UG fields and consolidated marksheets are cleanly wiped."),
        ("Marks Scale Change", "Changing scale from Percentage to CGPA did not re-evaluate marks until submit.", "Immediately triggers revalidation across all semesters on dropdown change."),
        ("Field Label State", "Labels remained plain grey when inputs were empty or failed validation.", "Labels highlight in bold red (text-danger-700) with an 'Invalid' badge."),
        ("Error Clearing", "Label errors did not react until full form re-submission.", "Highlighting clears automatically in real time as soon as valid input is entered."),
        ("Saved Drafts", "Drafts stored orphaned fields that blocked student upon resuming.", "Drafts are sanitized during merge to strip obsolete conditional data."),
        ("Vitest Test Suite", "5 test failures due to DOM duplicate errors and stale state.", "All 44 tests passing cleanly (100% pass rate in CI/local)."),
        ("TypeScript Status", "Unresolved types in test stubs.", "Zero errors (tsc -b --noEmit exits with code 0).")
    ]

    t = doc.add_table(rows=len(table_data), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)

    col_widths = [Inches(1.8), Inches(2.5), Inches(2.5)]

    for r_idx, row in enumerate(table_data):
        is_hdr = (r_idx == 0)
        bg = "3D3777" if is_hdr else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")

        for c_idx, text in enumerate(row):
            cell = t.rows[r_idx].cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1

            run = p.add_run(text)
            run.font.name = "Segoe UI"
            if is_hdr:
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.size = Pt(8.5)
                if c_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = PRIMARY
                else:
                    run.font.color.rgb = TEXT

    doc.save(str(output_path))
    print(f"Successfully generated docx at {output_path}")

if __name__ == "__main__":
    out_file = Path("docs/specs/2026-09-11-srf-dropdown-sync-and-label-validation-fixes.docx")
    out_file.parent.mkdir(parents=True, exist_ok=True)
    build_docx(out_file)
