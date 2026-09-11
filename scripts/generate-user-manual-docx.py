import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout(doc, text, callout_type="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    border_color = "3D3777"
    bg_color = "F0EEF8"
    icon_prefix = "📌 NOTE: "
    if callout_type == "WARNING":
        border_color = "DD4820"
        bg_color = "FFF0EC"
        icon_prefix = "⚠️ WARNING: "
    elif callout_type == "IMPORTANT":
        border_color = "FF7200"
        bg_color = "FFF5EB"
        icon_prefix = "⭐ IMPORTANT: "
    elif callout_type == "RULE":
        border_color = "3D3777"
        bg_color = "ECF1F0"
        icon_prefix = "⚖️ POLICY RULE: "
        
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_prefix = p.add_run(icon_prefix)
    r_prefix.bold = True
    r_prefix.font.size = Pt(10)
    r_prefix.font.name = "Calibri"
    r_prefix.font.color.rgb = RGBColor(0x3D, 0x37, 0x77) if callout_type != "WARNING" else RGBColor(0xDD, 0x48, 0x20)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(10)
    r_text.font.name = "Calibri"
    r_text.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def parse_markdown_to_docx(md_path: Path, docx_path: Path):
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("FACE Prep Campus — Placement Management System (PMS) | User Manual")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential — For Internal and Partner Campus Use Only")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # Style defaults
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(10.5)
    normal_font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_table = False
    table_rows = []
    
    while i < len(lines):
        line = lines[i].strip()
        
        # Detect table
        if line.startswith('|') and line.endswith('|'):
            table_rows.append([c.strip() for c in line[1:-1].split('|')])
            in_table = True
            i += 1
            continue
        elif in_table:
            # End of table, render it
            if table_rows:
                # Filter out separator row (e.g., |---|:--:|)
                filtered_rows = [r for r in table_rows if not all(re.match(r'^:?-+:?$', cell) for cell in r if cell)]
                if filtered_rows:
                    num_cols = max(len(r) for r in filtered_rows)
                    tbl = doc.add_table(rows=len(filtered_rows), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(tbl)
                    
                    for r_idx, row_data in enumerate(filtered_rows):
                        is_header = (r_idx == 0)
                        for c_idx in range(num_cols):
                            cell = tbl.cell(r_idx, c_idx)
                            val = row_data[c_idx] if c_idx < len(row_data) else ""
                            cell_p = cell.paragraphs[0]
                            cell_p.paragraph_format.space_before = Pt(3)
                            cell_p.paragraph_format.space_after = Pt(3)
                            cell_p.paragraph_format.line_spacing = 1.15
                            
                            run = cell_p.add_run(val)
                            run.font.name = 'Calibri'
                            run.font.size = Pt(9.5)
                            
                            if is_header:
                                set_cell_background(cell, "3D3777")
                                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                                run.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            else:
                                bg = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
                                set_cell_background(cell, bg)
                                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                                
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_before = Pt(0)
                    p_sp.paragraph_format.space_after = Pt(4)
            table_rows = []
            in_table = False

        if not line:
            i += 1
            continue

        # Markdown horizontal rule
        if re.match(r'^-{3,}$', line):
            i += 1
            continue

        # Title (# )
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(title_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(0x3D, 0x37, 0x77)
            i += 1
            continue

        # Subtitle (## Comprehensive User Manual...)
        if line.startswith('## ') and ('Comprehensive' in line or 'User Manual' in line):
            subtitle_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(subtitle_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # Heading 1 (## Chapter...)
        if line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(0x3D, 0x37, 0x77)
            i += 1
            continue

        # Heading 2 (### Subchapter...)
        if line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x4A, 0x42, 0x8C)
            i += 1
            continue

        # Heading 3 (#### Q1...)
        if line.startswith('#### '):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            i += 1
            continue

        # Warnings / Callouts
        if line.startswith('*Warning:*') or line.startswith('*Critical Rule:*') or 'Warning:' in line:
            clean_text = line.replace('*Warning:*', '').replace('*Critical Rule:*', '').strip()
            add_callout(doc, clean_text, callout_type="WARNING")
            i += 1
            continue
            
        if line.startswith('*Note:*') or 'Notice:' in line:
            clean_text = line.replace('*Note:*', '').replace('*Notice:*', '').strip()
            add_callout(doc, clean_text, callout_type="NOTE")
            i += 1
            continue

        # Bullet list (- or *)
        if line.startswith('- ') or line.startswith('* '):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            
            # Format bold within bullets
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Numbered list (e.g., 1. , 2. )
        m_num = re.match(r'^(\d+)\.\s+(.*)$', line)
        if m_num:
            num_str = m_num.group(1)
            item_text = m_num.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)
        i += 1

    # Render any remaining table at the end
    if table_rows:
        filtered_rows = [r for r in table_rows if not all(re.match(r'^:?-+:?$', cell) for cell in r if cell)]
        if filtered_rows:
            num_cols = max(len(r) for r in filtered_rows)
            tbl = doc.add_table(rows=len(filtered_rows), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl)
            for r_idx, row_data in enumerate(filtered_rows):
                for c_idx in range(num_cols):
                    cell = tbl.cell(r_idx, c_idx)
                    cell.paragraphs[0].text = row_data[c_idx] if c_idx < len(row_data) else ""

    doc.save(docx_path)
    print(f"Successfully generated docx at {docx_path}")

if __name__ == '__main__':
    md = Path("docs/USER_MANUAL.md")
    out = Path("docs/USER_MANUAL.docx")
    parse_markdown_to_docx(md, out)
